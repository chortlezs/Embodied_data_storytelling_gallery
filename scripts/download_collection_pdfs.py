#!/usr/bin/env python3
import json
import re
import time
import unicodedata
from difflib import SequenceMatcher
from html import unescape
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple
from urllib.parse import quote, urljoin

import requests
from openpyxl import load_workbook
from requests.exceptions import ChunkedEncodingError


ROOT = Path("/Users/bytedance/Desktop/tvcg案例库")
SOURCE_XLSX = ROOT / "collection.xlsx"
OUTPUT_DIR = ROOT / "paper_pdfs"
REPORT_PATH = ROOT / "paper_download_report.json"
TIMEOUT = (5, 12)
USER_AGENT = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0 Safari/537.36"
HEADERS = {"User-Agent": USER_AGENT, "Accept": "*/*"}


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = unescape(text)
    text = text.replace("’", "'").replace("`", "'")
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def similarity(a: str, b: str) -> float:
    return SequenceMatcher(None, normalize_text(a), normalize_text(b)).ratio()


def slugify(text: str, max_len: int = 120) -> str:
    text = normalize_text(text)
    text = text.replace(" ", "_").strip("_")
    text = re.sub(r"_+", "_", text)
    if not text:
        return "paper"
    return text[:max_len].rstrip("_")


def iter_titles() -> Iterable[str]:
    wb = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx < 4:
            continue
        title = row[2]
        if isinstance(title, str) and title.strip():
            yield title.strip()


def get_json(session: requests.Session, url: str) -> Dict:
    response = session.get(url, timeout=TIMEOUT, headers=HEADERS)
    response.raise_for_status()
    return response.json()


def get_best_openalex_match(session: requests.Session, title: str) -> Tuple[Optional[Dict], float]:
    search_url = f"https://api.openalex.org/works?search={quote(title)}&per-page=10"
    data = get_json(session, search_url)
    best = None
    best_score = 0.0
    for item in data.get("results", []):
        name = item.get("display_name") or ""
        score = similarity(title, name)
        if score > best_score:
            best = item
            best_score = score
    return best, best_score


def collect_candidate_urls(item: Dict) -> List[str]:
    urls: List[str] = []
    seen: Set[str] = set()

    def add(url: Optional[str]) -> None:
        if not url:
            return
        url = url.strip()
        if not url or url in seen:
            return
        seen.add(url)
        urls.append(url)

    add((item.get("best_oa_location") or {}).get("pdf_url"))
    add((item.get("best_oa_location") or {}).get("landing_page_url"))
    add((item.get("primary_location") or {}).get("pdf_url"))
    add((item.get("primary_location") or {}).get("landing_page_url"))

    for location in item.get("locations") or []:
        add(location.get("pdf_url"))
        add(location.get("landing_page_url"))

    open_access = item.get("open_access") or {}
    add(open_access.get("oa_url"))
    return urls[:8]


def is_pdf_response(response: requests.Response) -> bool:
    content_type = (response.headers.get("content-type") or "").lower()
    return "application/pdf" in content_type or response.url.lower().endswith(".pdf")


def looks_like_pdf_url(url: str) -> bool:
    lowered = url.lower()
    return ".pdf" in lowered or "/pdf" in lowered


def extract_pdf_links(html: str, base_url: str) -> List[str]:
    matches: List[str] = []
    patterns = [
        r'href=["\']([^"\']+)["\']',
        r'src=["\']([^"\']+)["\']',
    ]
    for pattern in patterns:
        for match in re.findall(pattern, html, flags=re.IGNORECASE):
            resolved = urljoin(base_url, unescape(match))
            if looks_like_pdf_url(resolved):
                matches.append(resolved)

    for match in re.findall(r'https?://[^"\')\s>]+', html, flags=re.IGNORECASE):
        resolved = unescape(match)
        if looks_like_pdf_url(resolved):
            matches.append(resolved)

    # Common meta tags for PDF/Scholarly landing pages.
    for pattern in [
        r'citation_pdf_url["\']?\s+content=["\']([^"\']+)["\']',
        r'content=["\']([^"\']+)["\']\s+name=["\']citation_pdf_url["\']',
    ]:
        for match in re.findall(pattern, html, flags=re.IGNORECASE):
            matches.append(urljoin(base_url, unescape(match)))

    deduped: List[str] = []
    seen: Set[str] = set()
    for url in matches:
        if url not in seen:
            deduped.append(url)
            seen.add(url)
    return deduped[:10]


def try_download_pdf(session: requests.Session, url: str, destination: Path, depth: int = 0) -> Tuple[bool, str]:
    try:
        response = session.get(url, timeout=TIMEOUT, headers=HEADERS, allow_redirects=True, stream=True)
        response.raise_for_status()
    except Exception as exc:
        return False, f"request_failed: {exc}"

    final_url = response.url
    if is_pdf_response(response):
        destination.parent.mkdir(parents=True, exist_ok=True)
        try:
            with destination.open("wb") as fh:
                for chunk in response.iter_content(chunk_size=65536):
                    if chunk:
                        fh.write(chunk)
        except (ChunkedEncodingError, OSError) as exc:
            destination.unlink(missing_ok=True)
            return False, f"stream_failed: {exc}"
        if destination.stat().st_size < 1024:
            destination.unlink(missing_ok=True)
            return False, "download_too_small"
        return True, final_url

    content_type = (response.headers.get("content-type") or "").lower()
    if "html" not in content_type:
        return False, f"not_pdf_content_type: {content_type or 'unknown'}"

    try:
        html = response.text
    except Exception as exc:
        return False, f"html_read_failed: {exc}"

    if depth >= 1:
        return False, "no_pdf_link_found"

    for candidate in extract_pdf_links(html, final_url):
        if candidate == url:
            continue
        ok, detail = try_download_pdf(session, candidate, destination, depth=depth + 1)
        if ok:
            return True, detail

    return False, "no_pdf_link_found"


def main() -> int:
    OUTPUT_DIR.mkdir(exist_ok=True)
    titles = list(iter_titles())
    session = requests.Session()
    session.headers.update(HEADERS)

    results: List[Dict] = []
    downloaded = 0

    for index, title in enumerate(titles, start=1):
        print(f"[{index}/{len(titles)}] {title}", flush=True)
        record: Dict = {
            "title": title,
            "status": "not_found",
            "match_score": 0.0,
            "matched_title": None,
            "doi": None,
            "source_url": None,
            "file_path": None,
            "notes": [],
        }

        try:
            item, score = get_best_openalex_match(session, title)
        except Exception as exc:
            record["status"] = "search_error"
            record["notes"].append(str(exc))
            results.append(record)
            continue

        record["match_score"] = round(score, 4)
        if not item or score < 0.72:
            record["notes"].append("no_confident_match_in_openalex")
            results.append(record)
            continue

        record["matched_title"] = item.get("display_name")
        record["doi"] = item.get("doi")

        target = OUTPUT_DIR / f"{slugify(title)}.pdf"
        if target.exists() and target.stat().st_size > 1024:
            record["status"] = "downloaded_existing"
            record["file_path"] = str(target)
            downloaded += 1
            results.append(record)
            continue

        candidate_urls = collect_candidate_urls(item)
        if not candidate_urls:
            record["status"] = "metadata_only"
            record["notes"].append("matched_but_no_open_url")
            results.append(record)
            continue

        for url in candidate_urls:
            ok, detail = try_download_pdf(session, url, target)
            if ok:
                record["status"] = "downloaded"
                record["source_url"] = detail
                record["file_path"] = str(target)
                downloaded += 1
                break
            record["notes"].append(f"{url} -> {detail}")

        if record["status"] != "downloaded":
            record["status"] = "open_version_not_downloaded"

        results.append(record)
        time.sleep(0.2)

    report = {
        "source": str(SOURCE_XLSX),
        "output_dir": str(OUTPUT_DIR),
        "downloaded_count": downloaded,
        "total_titles": len(titles),
        "results": results,
    }
    REPORT_PATH.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nDownloaded {downloaded} / {len(titles)} PDFs")
    print(f"PDF directory: {OUTPUT_DIR}")
    print(f"Report: {REPORT_PATH}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
